from django.shortcuts import render, redirect, get_object_or_404
from django.views.generic import TemplateView, CreateView, UpdateView, DetailView
from .models import Categoria, Producto, Compra, DetalleCompra, Venta, DetalleVenta, Proveedor, Cliente
from .forms import CategoriaForm, ProductoForm, CompraForm, VentaForm, ProveedorForm, ClienteForm, CarritoProductoForm
from django.urls import reverse_lazy, reverse
from django.contrib import messages
from django.http import HttpResponse
from datetime import datetime, date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
# Create your views here.

def ListarCategorias(request):
    categorias = Categoria.objects.filter(activo=True)
    return render(request, 'categorias.html', {'categorias': categorias})

def ListarProveedores(request):
    proveedores = Proveedor.objects.filter(activo=True)
    return render(request, 'proveedores.html', {'proveedores': proveedores})

class CrearProveedor(CreateView):
    template_name = 'crear_proveedor.html'
    form_class = ProveedorForm
    success_url = reverse_lazy('administradores:proveedoresapp')

class EditarProveedor(UpdateView):
    template_name = 'editar_proveedor.html'
    form_class = ProveedorForm
    model = Proveedor
    success_url = reverse_lazy('administradores:proveedoresapp')

def ToggleProveedorActivo(request, pk):
    proveedor = get_object_or_404(Proveedor, pk=pk)
    if request.method == 'POST':
        proveedor.activo = not proveedor.activo
        proveedor.save()
        messages.success(request, 'El proveedor fue ocultado correctamente.' if not proveedor.activo else 'El proveedor fue reactivado correctamente.')
        return redirect('administradores:proveedoresapp')

def ListarClientes(request):
    clientes = Cliente.objects.all()
    return render(request, 'clientes.html', {'clientes': clientes})

class CrearCliente(CreateView):
    template_name = 'crear_cliente.html'
    form_class = ClienteForm
    success_url = reverse_lazy('administradores:clientesapp')

class EditarCliente(UpdateView):
    template_name = 'editar_cliente.html'
    form_class = ClienteForm
    model = Cliente
    success_url = reverse_lazy('administradores:clientesapp')

def ListarProductosCategoria(request, cid):
    categoria = get_object_or_404(Categoria, pk=cid)
    productos = Producto.objects.filter(id_categoria=categoria, activo=True)
    return render(request, 'productos.html', {'productos': productos, 'categoria': categoria})

class CrearCategoria(CreateView):
    template_name = 'crear_categoria.html'
    form_class = CategoriaForm
    success_url = reverse_lazy('administradores:productosapp')

class EditarCategoria(UpdateView):
    template_name = 'editar_categoria.html'
    form_class = CategoriaForm
    model = Categoria
    success_url = reverse_lazy('administradores:productosapp')

class CrearProducto(CreateView):
    template_name = 'crear_producto.html'
    form_class = ProductoForm

    def get_initial(self):
        initial = super().get_initial()
        initial['id_categoria'] = self.kwargs.get('cid')
        return initial

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['categoria'] = get_object_or_404(Categoria, pk=self.kwargs.get('cid'))
        return context

    def get_success_url(self):
        return reverse('administradores:productoscat', args=[self.kwargs.get('cid')])

    def form_valid(self, form):
        form.instance.cantidad = 0
        return super().form_valid(form)

class EditarProducto(UpdateView):
    template_name = 'editar_producto.html'
    form_class = ProductoForm
    model = Producto

    def get_success_url(self):
        return reverse('administradores:productoscat', args=[self.object.id_categoria_id])

class VerProducto(DetailView):
    template_name = 'ver_producto.html'
    model = Producto

def ListarCompras(request):
    compras = Compra.objects.filter(activo=True)
    return render(request, 'compras.html', {'compras': compras})

def ListarVentas(request):
    ventas = Venta.objects.filter(activo=True)
    return render(request, 'ventas.html', {'ventas': ventas})


def _cart_items(carrito, price_attr='precio_compra'):
    items = []
    total = 0
    for pid, cantidad in carrito.items():
        producto = get_object_or_404(Producto, pk=pid)
        precio = getattr(producto, price_attr) or 0
        subtotal = precio * cantidad
        total += subtotal
        items.append({
            'id': producto.id,
            'producto': producto,
            'nombre': producto.nombre,
            'cantidad': cantidad,
            'precio_unitario': precio,
            'subtotal': subtotal,
        })
    return items, total


def CrearCompra(request):
    pendiente = request.session.get('compra_pendiente', {})
    initial = {}
    if pendiente:
        initial['fecha'] = pendiente.get('fecha')
        initial['id_proveedor'] = pendiente.get('id_proveedor')
        initial['id_usuario'] = pendiente.get('id_usuario')

    if request.method == 'POST':
        form = CompraForm(request.POST)
        accion = request.POST.get('accion')
        if form.is_valid():
            fechaseleccionada = form.cleaned_data.get('fecha')
            request.session['compra_pendiente'] = {
                'fecha': fechaseleccionada.isoformat() if fechaseleccionada else None,
                'id_proveedor': form.cleaned_data['id_proveedor'].pk,
                'id_usuario': form.cleaned_data['id_usuario'].pk if form.cleaned_data['id_usuario'] else None,
            }
            carrito = request.session.get('carrito_compra', {})
            if accion == 'ver_carrito':
                return redirect('administradores:vercarritoapp')
            elif accion == 'confirmar':
                if not carrito:
                    messages.error(request, 'El carrito está vacío. Agregue productos antes de confirmar la compra.')
                else:
                    items, total = _cart_items(carrito)
                    compra = Compra.objects.create(
                        fecha=fechaseleccionada,
                        total=total,
                        id_proveedor=form.cleaned_data['id_proveedor'],
                        id_usuario=form.cleaned_data['id_usuario'],
                    )
                    for item in items:
                        DetalleCompra.objects.create(
                            id_compra=compra,
                            id_producto=item['producto'],
                            cantidad=item['cantidad'],
                            precio_unitario=item['precio_unitario'],
                            subtotal=item['subtotal'],
                        )
                        producto = item['producto']
                        producto.cantidad = (producto.cantidad or 0) + item['cantidad']
                        producto.save()
                    request.session.pop('carrito_compra', None)
                    request.session.pop('compra_pendiente', None)
                    messages.success(request, 'La compra fue registrada correctamente.')
                    return redirect('administradores:comprasapp')
    else:
        form = CompraForm(initial=initial)

    return render(request, 'crear_compra.html', {'form': form})


def VerCarrito(request):
    carrito = request.session.get('carrito_compra', {})
    items, total = _cart_items(carrito)
    return render(request, 'carrito_compra.html', {'items': items, 'total': total})


def AgregarProductoCarrito(request):
    context = {}
    if request.method == 'POST':
        form = CarritoProductoForm(request.POST)
        if form.is_valid():
            producto = form.cleaned_data['producto']
            cantidad = form.cleaned_data['cantidad']
            carrito = request.session.get('carrito_compra', {})
            carrito[str(producto.pk)] = carrito.get(str(producto.pk), 0) + cantidad
            request.session['carrito_compra'] = carrito
            messages.success(request, 'Producto agregado al carrito.')
            return redirect('administradores:vercarritoapp')
    else:
        form = CarritoProductoForm()
    context['form'] = form
    return render(request, 'agregar_producto_carrito.html', context)


def QuitarProductoCarrito(request, pk):
    if request.method == 'POST':
        carrito = request.session.get('carrito_compra', {})
        carrito.pop(str(pk), None)
        request.session['carrito_compra'] = carrito
        messages.success(request, 'Producto eliminado del carrito.')
    return redirect('administradores:vercarritoapp')

class EditarCompra(UpdateView):
    template_name = 'editar_compra.html'
    form_class = CompraForm
    model = Compra
    success_url = reverse_lazy('administradores:comprasapp')

class VerCompra(DetailView):
    template_name = 'ver_compra.html'
    model = Compra

def CrearVenta(request):
    pendiente = request.session.get('venta_pendiente', {})
    initial = {}
    if pendiente:
        initial['fecha'] = pendiente.get('fecha')
        initial['id_cliente'] = pendiente.get('id_cliente')

    if request.method == 'POST':
        form = VentaForm(request.POST)
        accion = request.POST.get('accion')
        if form.is_valid():
            fechaseleccionada = form.cleaned_data.get('fecha')
            request.session['venta_pendiente'] = {
                'fecha': fechaseleccionada.isoformat() if fechaseleccionada else None,
                'id_cliente': form.cleaned_data['id_cliente'].pk,
            }
            carrito = request.session.get('carrito_venta', {})
            if accion == 'ver_carrito':
                return redirect('administradores:vercarritoventaapp')
            elif accion == 'confirmar':
                if not carrito:
                    messages.error(request, 'El carrito está vacío. Agregue productos antes de confirmar la venta.')
                else:
                    items, total = _cart_items(carrito, 'precio_venta')
                    venta = Venta.objects.create(
                        fecha=fechaseleccionada,
                        total=total,
                        id_cliente=form.cleaned_data['id_cliente'],
                        id_usuario=form.cleaned_data.get('id_usuario'),
                    )
                    for item in items:
                        DetalleVenta.objects.create(
                            id_venta=venta,
                            id_producto=item['producto'],
                            cantidad=item['cantidad'],
                            precio_unitario=item['precio_unitario'],
                            subtotal=item['subtotal'],
                        )
                        producto = item['producto']
                        producto.cantidad = (producto.cantidad or 0) - item['cantidad']
                        producto.save()
                    request.session.pop('carrito_venta', None)
                    request.session.pop('venta_pendiente', None)
                    messages.success(request, 'La venta fue registrada correctamente.')
                    return redirect('administradores:ventasapp')
    else:
        form = VentaForm(initial=initial)

    return render(request, 'crear_venta.html', {'form': form})


def VerCarritoVenta(request):
    carrito = request.session.get('carrito_venta', {})
    items, total = _cart_items(carrito, 'precio_venta')
    return render(request, 'carrito_venta.html', {'items': items, 'total': total})


def AgregarProductoVenta(request):
    context = {}
    if request.method == 'POST':
        form = CarritoProductoForm(request.POST)
        if form.is_valid():
            producto = form.cleaned_data['producto']
            cantidad = form.cleaned_data['cantidad']
            if producto.cantidad is None or producto.cantidad <= 0:
                messages.error(request, 'El producto "%s" ha sido agotado y no puede agregarse al carrito.' % producto.nombre)
            else:
                carrito = request.session.get('carrito_venta', {})
                actual = carrito.get(str(producto.pk), 0)
                if actual + cantidad > producto.cantidad:
                    messages.error(request, 'No hay suficiente existencia de "%s". Disponibles: %s unidades.' % (producto.nombre, producto.cantidad))
                else:
                    carrito[str(producto.pk)] = actual + cantidad
                    request.session['carrito_venta'] = carrito
                    messages.success(request, 'Producto agregado al carrito.')
                    return redirect('administradores:vercarritoventaapp')
    else:
        form = CarritoProductoForm()
    context['form'] = form
    return render(request, 'agregar_producto_venta.html', context)


def QuitarProductoVenta(request, pk):
    if request.method == 'POST':
        carrito = request.session.get('carrito_venta', {})
        carrito.pop(str(pk), None)
        request.session['carrito_venta'] = carrito
        messages.success(request, 'Producto eliminado del carrito.')
    return redirect('administradores:vercarritoventaapp')


class EditarVenta(UpdateView):
    template_name = 'editar_venta.html'
    form_class = VentaForm
    model = Venta
    success_url = reverse_lazy('administradores:ventasapp')

class VerVenta(DetailView):
    template_name = 'ver_venta.html'
    model = Venta

def _parse_fecha(valor):
    if not valor:
        return None
    try:
        return datetime.strptime(valor, '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None


def _reporte_productos(categoria_id):
    productos = Producto.objects.filter(activo=True).select_related('id_categoria')
    if categoria_id:
        productos = productos.filter(id_categoria_id=categoria_id)
    return productos.order_by('id_categoria_id', 'nombre')


def _reporte_compras(fecha_inicio, fecha_fin):
    compras = Compra.objects.filter(activo=True).select_related('id_proveedor')
    fi = _parse_fecha(fecha_inicio)
    ff = _parse_fecha(fecha_fin)
    if fi:
        compras = compras.filter(fecha__gte=fi)
    if ff:
        compras = compras.filter(fecha__lte=ff)
    return compras.order_by('fecha', 'id').prefetch_related('detallecompra_set__id_producto')


def _reporte_ventas(fecha_inicio, fecha_fin):
    ventas = Venta.objects.filter(activo=True).select_related('id_cliente')
    fi = _parse_fecha(fecha_inicio)
    ff = _parse_fecha(fecha_fin)
    if fi:
        ventas = ventas.filter(fecha__gte=fi)
    if ff:
        ventas = ventas.filter(fecha__lte=ff)
    return ventas.order_by('fecha', 'id').prefetch_related('detalleventa_set__id_producto')


def _celda_pdf(valor):
    if valor is None:
        return ''
    if isinstance(valor, date):
        return valor.strftime('%d/%m/%Y')
    if isinstance(valor, float):
        return 'Q%.2f' % valor
    return str(valor)


def _txt_rango(fecha_inicio, fecha_fin):
    partes = []
    if fecha_inicio:
        partes.append('desde %s' % fecha_inicio)
    if fecha_fin:
        partes.append('hasta %s' % fecha_fin)
    return ' (%s)' % ' '.join(partes) if partes else ''


def _datos_exportacion(tipo, post):
    if tipo == 'productos':
        categoria_id = post.get('id_categoria')
        productos = _reporte_productos(categoria_id)
        encabezados = ['ID', 'Nombre', 'Precio Compra', 'Precio Venta', 'Existencia', 'Fecha Vencimiento']
        anchos = [40, 130, 80, 80, 70, 130]
        filas = []
        for p in productos:
            fila = [p.id, p.nombre, p.precio_compra, p.precio_venta, p.cantidad, p.fecha_vencimiento]
            if not categoria_id:
                encabezados.append('Categoría')
                fila.append(p.id_categoria.nombre)
            filas.append(fila)
        titulo = 'Reporte de Productos'
        if categoria_id:
            titulo += ' - %s' % Categoria.objects.get(pk=categoria_id).nombre
        return titulo, encabezados, filas, anchos, [2, 3]
    if tipo == 'compras':
        compras = _reporte_compras(post.get('fecha_inicio'), post.get('fecha_fin'))
        titulo = 'Reporte de Compras' + _txt_rango(post.get('fecha_inicio'), post.get('fecha_fin'))
        encabezados = ['ID', 'Fecha', 'Proveedor', 'Total', 'Producto', 'Cantidad', 'Precio Unitario', 'Subtotal']
        filas = []
        for c in compras:
            detalles = c.detallecompra_set.all()
            if not detalles:
                filas.append([c.id, c.fecha, c.id_proveedor.nombre, c.total, '-', '', '', ''])
            for d in detalles:
                filas.append([c.id, c.fecha, c.id_proveedor.nombre, c.total, d.id_producto.nombre, d.cantidad, d.precio_unitario, d.subtotal])
        return titulo, encabezados, filas, [46, 70, 125, 55, 160, 55, 70, 70], [3, 6, 7]
    if tipo == 'ventas':
        ventas = _reporte_ventas(post.get('fecha_inicio'), post.get('fecha_fin'))
        titulo = 'Reporte de Ventas' + _txt_rango(post.get('fecha_inicio'), post.get('fecha_fin'))
        encabezados = ['ID', 'Fecha', 'Cliente', 'Total', 'Producto', 'Cantidad', 'Precio Unitario', 'Subtotal']
        filas = []
        for v in ventas:
            detalles = v.detalleventa_set.all()
            if not detalles:
                filas.append([v.id, v.fecha, str(v.id_cliente), v.total, '-', '', '', ''])
            for d in detalles:
                filas.append([v.id, v.fecha, str(v.id_cliente), v.total, d.id_producto.nombre, d.cantidad, d.precio_unitario, d.subtotal])
        return titulo, encabezados, filas, [46, 70, 125, 55, 160, 55, 70, 70], [3, 6, 7]
    if tipo == 'clientes':
        clientes = Cliente.objects.order_by('apellido', 'nombre')
        titulo = 'Reporte de Clientes'
        encabezados = ['ID', 'Nombre', 'Apellido']
        filas = [[c.id, c.nombre, c.apellido] for c in clientes]
        return titulo, encabezados, filas, [60, 150, 150], []
    if tipo == 'proveedores':
        proveedores = Proveedor.objects.filter(activo=True).order_by('nombre')
        titulo = 'Reporte de Proveedores'
        encabezados = ['ID', 'Nombre', 'Teléfono', 'Dirección']
        filas = [[p.id, p.nombre, p.telefono or '', p.direccion or ''] for p in proveedores]
        return titulo, encabezados, filas, [55, 160, 120, 130], []
    return None


def _responder_pdf(titulo, encabezados, filas, anchos):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        title=titulo,
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20,
    )
    elementos = [Paragraph(titulo, getSampleStyleSheet()['Title']), Spacer(1, 14)]
    data = [encabezados] + [[_celda_pdf(v) for v in fila] for fila in filas]
    tabla = Table(data, colWidths=anchos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#343a40')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f2f2f2')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    respuesta = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    return respuesta


def _responder_excel(titulo, encabezados, filas, money_cols):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte'
    ws.append([titulo])
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(encabezados))
    ws.append([])
    ws.append(encabezados)
    for celda in ws[ws.max_row]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill('solid', fgColor='D9E1F2')
    for fila in filas:
        ws.append(['' if v is None else v for v in fila])
    for fila_idx in range(5, ws.max_row + 1):
        for col_idx in money_cols:
            celda = ws.cell(row=fila_idx, column=col_idx)
            if isinstance(celda.value, (int, float)):
                celda.number_format = '"Q"#,##0.00'
    for i in range(1, len(encabezados) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    from django.http import HttpResponse
    respuesta = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    wb.save(respuesta)
    return respuesta


def _exportar_reporte(tipo, post, formato):
    datos = _datos_exportacion(tipo, post)
    if not datos:
        return None
    titulo, encabezados, filas, anchos, money_cols = datos
    slug = tipo or 'reporte'
    ext = 'pdf' if formato == 'pdf' else 'xlsx'
    if formato == 'pdf':
        respuesta = _responder_pdf(titulo, encabezados, filas, anchos)
    else:
        respuesta = _responder_excel(titulo, encabezados, filas, money_cols)
    respuesta['Content-Disposition'] = 'attachment; filename="reporte_%s.%s"' % (slug, ext)
    return respuesta


def Reportes(request):
    context = {
        'tipos': [
            ('productos', 'Productos'),
            ('compras', 'Compras'),
            ('ventas', 'Ventas'),
            ('clientes', 'Clientes'),
            ('proveedores', 'Proveedores'),
        ],
        'categorias': Categoria.objects.filter(activo=True),
    }
    if request.method == 'POST':
        tipo = request.POST.get('tipo')
        formato = request.POST.get('formato')
        context['tipo'] = tipo
        context['tipo_seleccionado'] = tipo
        if tipo == 'productos':
            context['categoria_id'] = request.POST.get('id_categoria')
            if context['categoria_id']:
                context['categoria_nombre'] = Categoria.objects.get(pk=context['categoria_id']).nombre
        elif tipo in ('compras', 'ventas'):
            context['fecha_inicio'] = request.POST.get('fecha_inicio')
            context['fecha_fin'] = request.POST.get('fecha_fin')
        if formato in ('pdf', 'excel'):
            respuesta = _exportar_reporte(tipo, request.POST, formato)
            if respuesta:
                return respuesta
            messages.error(request, 'Seleccione un tipo de reporte para exportar.')
        if tipo == 'productos':
            context['productos'] = _reporte_productos(context.get('categoria_id'))
        elif tipo == 'compras':
            context['compras'] = _reporte_compras(context.get('fecha_inicio'), context.get('fecha_fin'))
        elif tipo == 'ventas':
            context['ventas'] = _reporte_ventas(context.get('fecha_inicio'), context.get('fecha_fin'))
        elif tipo == 'clientes':
            context['clientes'] = Cliente.objects.order_by('apellido', 'nombre')
        elif tipo == 'proveedores':
            context['proveedores'] = Proveedor.objects.filter(activo=True).order_by('nombre')
    return render(request, 'reportes.html', context)

class LoginView(TemplateView):
    template_name = 'login.html'

def ToggleProductoActivo(request, pk):
    producto = get_object_or_404(Producto, pk=pk)
    if request.method == 'POST':
        producto.activo = not producto.activo
        producto.save()
        return redirect('administradores:productoscat', cid=producto.id_categoria_id)

def ToggleCategoriaActivo(request, pk):
    categoria = get_object_or_404(Categoria, pk=pk)
    if request.method == 'POST':
        categoria.activo = not categoria.activo
        categoria.save()
        messages.success(request, 'La categoría fue ocultada correctamente.' if not categoria.activo else 'La categoría fue reactivada correctamente.')
        return redirect('administradores:productosapp')

def ToggleCompraActivo(request, pk):
    compra = get_object_or_404(Compra, pk=pk)
    if request.method == 'POST':
        compra.activo = not compra.activo
        compra.save()
        return redirect('administradores:comprasapp')

def ToggleVentaActivo(request, pk):
    venta = get_object_or_404(Venta, pk=pk)
    if request.method == 'POST':
        venta.activo = not venta.activo
        venta.save()
        return redirect('administradores:ventasapp')
